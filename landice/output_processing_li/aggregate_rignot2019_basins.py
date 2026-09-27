#!/usr/bin/env python3
"""Aggregate the main table in pnas.1812883116.sd01-CLEAN.xlsx.

Requires openpyxl: python -m pip install openpyxl
Usage: python aggregate_basins.py INPUT.xlsx -o basin_summary.txt

Groups by column B (the two-line 'Basin Name' header). All fluxes are
in Gt/yr. RMS means sqrt(sum(sigma**2)/N), NOT root-sum-of-squares.
D is sum over glaciers of each glacier's arithmetic mean over the inclusive
--start-year / --end-year range (defaults: 2009 / 2017).
Leading whitespace in column A identifies subdivisions to exclude.
Font color is ignored; the workbook legend about grey text is incorrect.
Reads cached Excel formula results; missing/non-numeric data raise errors.
Outputs Python dictionary entries with input, outflow, and net [value, sigma].
Original basin RMS uncertainties are retained. For merged E-F and J-K,
original basin uncertainties are combined in quadrature, as in the example.
Net = input - outflow; net sigma assumes independent input/outflow errors.
Output is rounded to one decimal place only after all calculations.

Output is produced in a format that can be pasted directly into
plot_regionalStats.py.
"""

import argparse
import math
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl


def number(cell):
    value = cell.value
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f'{cell.coordinate}: expected a numeric value; got {value!r}. '
                         'If this cell contains a formula, recalculate and save in Excel first.')
    if not math.isfinite(value):
        raise ValueError(f'{cell.coordinate}: non-finite value {value!r}')
    return float(value)


def aggregate(path, sheet=None, verbose=False, start_year=2009, end_year=2017):
    if start_year > end_year:
        raise ValueError("Start year must be less than or equal to end year.")
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = workbook[sheet] if sheet else workbook.worksheets[0]
        if (ws['B1'].value, ws['B2'].value) != ('Basin', 'Name'):
            raise ValueError('Expected Basin / Name header in B1:B2.')
        # Annual discharge occupies K:AW in this workbook (1979-2017).
        year_columns = {}
        for column in range(11, 50):
            year = ws.cell(1, column).value
            if isinstance(year, (int, float)) and not isinstance(year, bool) and year == int(year):
                year = int(year)
                if year in year_columns:
                    raise ValueError(f'Duplicate discharge year header: {year}')
                year_columns[year] = column
        years = range(start_year, end_year + 1)
        missing = [year for year in years if year not in year_columns]
        if missing:
            raise ValueError(f'Discharge years not available in K:AW: {missing}')
        columns = [year_columns[year] for year in years]
        groups = defaultdict(list)
        excluded = 0
        for row in ws.iter_rows(min_row=5, max_col=49):
            # The first entirely empty A:I row ends the main data table.
            # Later rows contain supporting tables, not glacier records.
            if all(c.value is None for c in row[:9]):
                break
            basin = row[1].value
            if basin is None or not str(basin).strip():
                continue  # regional totals and TOTAL SURVEYED
            if not isinstance(row[0].value, str) or not row[0].value.strip():
                raise ValueError(f'A{row[0].row}: missing glacier name')
            if row[0].value != row[0].value.lstrip():
                excluded += 1
                if verbose:
                    print(f'Excluded row {row[0].row}: {row[0].value.strip()}', file=sys.stderr)
                continue
            sigma_smb, smb, sigma_d = (number(row[c - 1]) for c in (6, 7, 9))
            mean_d = math.fsum(number(row[c - 1]) for c in columns) / len(columns)
            groups[str(basin).strip()].append((smb, sigma_smb, mean_d, sigma_d))
        if not groups:
            raise ValueError('No glacier records found.')
        result = []
        for basin, values in groups.items():
            n = len(values)
            result.append({
                'Basin Name': basin,
                'n_glaciers': n,
                'SMB_sum_Gt_per_yr': math.fsum(v[0] for v in values),
                'sigma_SMB_RMS_Gt_per_yr': math.sqrt(math.fsum(v[1]**2 for v in values) / n),
                'D_mean_sum_Gt_per_yr': math.fsum(v[2] for v in values),
                'sigma_D_RMS_Gt_per_yr': math.sqrt(math.fsum(v[3]**2 for v in values) / n),
            })
        return result, excluded
    finally:
        workbook.close()


# Explicit mapping from workbook names to the requested 16 ISMIP6 basins.
# I"J is the workbook's spelling of I"-J.
ISMIP6_BASINS = {
    'ISMIP6BasinAAp': ("A-A'",),
    'ISMIP6BasinApB': ("A'-B",),
    'ISMIP6BasinBC': ('B-C',),
    'ISMIP6BasinCCp': ("C-C'",),
    'ISMIP6BasinCpD': ("C'-D",),
    'ISMIP6BasinDDp': ("D-D'",),
    'ISMIP6BasinDpE': ("D'-E",),
    'ISMIP6BasinEF': ("E-E'", "E'-F"),
    'ISMIP6BasinFG': ('F-G',),
    'ISMIP6BasinGH': ('G-H',),
    'ISMIP6BasinHHp': ("H-H'",),
    'ISMIP6BasinHpI': ("H'-I",),
    'ISMIP6BasinIIpp': ('I-I"',),
    'ISMIP6BasinIppJ': ('I"J',),
    'ISMIP6BasinJK': ('J-J"', 'J"-K'),
    'ISMIP6BasinKA': ('K-A',),
}


def ismip6_entries(result):
    by_name = {row['Basin Name']: row for row in result}
    expected = {name for names in ISMIP6_BASINS.values() for name in names}
    if set(by_name) != expected:
        raise ValueError(f'Basin mapping mismatch: missing={sorted(expected - set(by_name))}, '
                         f'unmapped={sorted(set(by_name) - expected)}')
    entries = {}
    for key, names in ISMIP6_BASINS.items():
        rows = [by_name[name] for name in names]
        smb = math.fsum(r['SMB_sum_Gt_per_yr'] for r in rows)
        discharge = math.fsum(r['D_mean_sum_Gt_per_yr'] for r in rows)
        sigma_smb = math.sqrt(math.fsum(r['sigma_SMB_RMS_Gt_per_yr']**2 for r in rows))
        sigma_d = math.sqrt(math.fsum(r['sigma_D_RMS_Gt_per_yr']**2 for r in rows))
        entries[key] = {'input': [smb, sigma_smb],
                        'outflow': [discharge, sigma_d],
                        'net': [smb - discharge, math.hypot(sigma_smb, sigma_d)]}
    return entries


def write_entries(stream, entries):
    # Dictionary fragment ready to paste between { and }, matching the example.
    for key, entry in entries.items():
        fields = ', '.join(f"{name!r}: [{values[0]:.1f}, {values[1]:.1f}]"
                           for name, values in entry.items())
        stream.write(f"                {key!r}: {{{fields}}},\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('input', type=Path, help='Input Excel workbook')
    parser.add_argument('-o', '--output', type=Path, default=Path('basin_summary.txt'))
    parser.add_argument('--start-year', type=int, default=2009,
                        help='First discharge year, inclusive (default: 2009)')
    parser.add_argument('--end-year', type=int, default=2017,
                        help='Last discharge year, inclusive (default: 2017)')
    parser.add_argument('--sheet', help='Worksheet name; defaults to first worksheet')
    parser.add_argument('-v', '--verbose', action='store_true', help='List excluded rows')
    args = parser.parse_args()
    if args.input.resolve() == args.output.resolve():
        parser.error('Output must differ from input.')
    try:
        result, excluded = aggregate(args.input, args.sheet, args.verbose,
                                     args.start_year, args.end_year)
        entries = ismip6_entries(result)
        with args.output.open('w', encoding='utf-8') as stream:
            write_entries(stream, entries)
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(1, f'Error: {exc}\n')
    print(f'Wrote {len(entries)} ISMIP6 basins to {args.output}; '
          f'included {sum(r["n_glaciers"] for r in result)} glacier rows; '
          f'excluded {excluded} indented rows; D period {args.start_year}-{args.end_year}.', file=sys.stderr)


if __name__ == '__main__':
    main()
